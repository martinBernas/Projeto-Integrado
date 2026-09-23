"""Generate the incremental import from the verified, versioned workbook."""
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import subprocess
import openpyxl

root = Path(__file__).resolve().parents[3]
source = subprocess.check_output(['git', 'show', 'HEAD:GeoGuaras.xlsx'], cwd=root)
assert sha256(source).hexdigest() == 'e74d1a6560a56ee1f7d3103df641199fbf308b2a5e735a46b6e0082bf93efcc4'
sheet = openpyxl.load_workbook(BytesIO(source), data_only=True)['Diario']
players = {'Luca': '8d899496-3cb0-41f6-b67d-4e57bb1185c4', 'Zade': '55a4aeaf-30bc-403a-a5e6-563328466313'}
scores = []
for row in range(2, 17):
    name = sheet.cell(row, 1).value
    if name not in players:
        continue
    for col in range(4, sheet.max_column + 1):
        day, value = sheet.cell(1, col).value, sheet.cell(row, col).value
        if isinstance(day, datetime) and datetime(2026, 9, 1) <= day <= datetime(2026, 9, 22) and value is not None:
            assert type(value) is int and 0 <= value <= 25000
            scores.append(f"('{name}', '{day.date()}', {value})")
assert len(scores) == 29
sql = """-- Luca e Zade: UUIDs e inicio em 01/09 confirmados pelo Dono do produto em 23/09.
-- Fonte: GeoGuaras.xlsx versionado, aba Diario, ate 22/09/2026.
-- SHA256: e74d1a6560a56ee1f7d3103df641199fbf308b2a5e735a46b6e0082bf93efcc4
-- 14 resultados de Luca e 15 de Zade. Nao inclui 23/09.
-- Executar integralmente no SQL Editor do Supabase como administrador.
-- Preserva a preparacao anterior: se penalidades estavam ativas, reativa ao final.
begin;
select id from public.tournaments where id='20260900-0000-4000-8000-000000000001' for update;
create temporary table new_players(name text primary key, player_id uuid unique not null) on commit drop;
insert into new_players values
('Luca', '8d899496-3cb0-41f6-b67d-4e57bb1185c4'),
('Zade', '55a4aeaf-30bc-403a-a5e6-563328466313');
create temporary table previous_history on commit drop as
select history_ready from public.tournaments where id = '20260900-0000-4000-8000-000000000001';
create temporary table new_scores(name text references new_players(name), game_day date, raw_score integer,
  primary key(name, game_day), check(raw_score between 0 and 25000)) on commit drop;
insert into new_scores values
""" + ',\n'.join(scores) + """;
do $$
declare p record; score_row record;
begin
  if not exists(select 1 from previous_history) then
    raise exception 'Torneio de setembro nao encontrado';
  end if;
  if exists(select 1 from new_players n left join auth.users u on u.id=n.player_id
    left join public.profiles pr on pr.id=u.id where u.id is null or pr.id is null) then
    raise exception 'Conta ou perfil ausente. Confira os UUIDs antes da carga';
  end if;
  for p in select * from new_players loop
    perform private.add_september_participant(p.player_id, p.name);
    update public.tournament_participants set eligible_from = date '2026-09-01'
      where tournament_id = '20260900-0000-4000-8000-000000000001' and player_id = p.player_id;
  end loop;
  for score_row in select n.player_id, s.game_day, s.raw_score from new_scores s join new_players n using(name)
    order by s.game_day, n.name loop
    perform private.load_september_score(score_row.player_id, score_row.game_day, score_row.raw_score);
  end loop;
  if exists(select 1 from new_scores s join new_players n using(name)
    left join public.personal_scores a on a.player_id=n.player_id and a.played_on=s.game_day
    where a.id is null or a.score <> s.raw_score) then
    raise exception 'Falha na conciliacao das 29 pontuacoes';
  end if;
  if (select history_ready from previous_history) then
    perform private.complete_september_history();
  end if;
end $$;
select n.name, n.player_id, p.eligible_from, count(s.game_day) as scores_imported
from new_players n join public.tournament_participants p on p.player_id=n.player_id
  and p.tournament_id='20260900-0000-4000-8000-000000000001'
join new_scores s using(name) group by n.name,n.player_id,p.eligible_from order by n.name;
select history_ready from public.tournaments where id='20260900-0000-4000-8000-000000000001';
commit;
"""
Path(__file__).with_name('september-new-players-2026-09-23.sql').write_text(sql, encoding='utf-8')
print('Luca: 14; Zade: 15. SQL incremental gerado.')
