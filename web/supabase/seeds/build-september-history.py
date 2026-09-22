"""Extract the confirmed column-C accounts; never edit the source workbook."""
from datetime import datetime
from hashlib import sha256
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[3]
source = root / 'GeoGuaras.xlsx'
sheet = openpyxl.load_workbook(source, data_only=True)['Diario']
cutoff = datetime(2026, 9, 22)
columns = [c for c in range(4, sheet.max_column + 1)
           if isinstance(sheet.cell(1, c).value, datetime)
           and datetime(2026, 9, 1) <= sheet.cell(1, c).value <= cutoff]
players, scores, pending = [], [], []
def quote(value):
    return "'" + str(value).replace("'", "''") + "'"
for row in range(2, 17):
    name, account = sheet.cell(row, 1).value, sheet.cell(row, 3).value
    if not account:
        pending.append(name)
        continue
    assert account == account.strip() and account not in [p[1] for p in players]
    players.append((name, account))
    for col in columns:
        value = sheet.cell(row, col).value
        if value is None:
            continue
        assert type(value) is int and 0 <= value <= 25000, sheet.cell(row, col).coordinate
        scores.append((account, sheet.cell(1, col).value.date().isoformat(), value))

sql = f"""-- GeoGuaras.xlsx / Diario, 01 a 22/09/2026.
-- SHA256: {sha256(source.read_bytes()).hexdigest()}
-- {len(players)} contas confirmadas na coluna C; {len(scores)} pontuacoes brutas.
-- Sem cadastro (nao incluidos): {', '.join(pending)}.
-- Execute integralmente no SQL Editor do Supabase, como administrador.
-- Requer as tres migracoes existentes e o seed september.sql ja aplicados.
-- Nao cria contas. Nao conclui o historico nem ativa penalidades por faltas.
-- O cadastro usa e-mail: identificadores C devem corresponder EXATAMENTE
-- a parte anterior ao @ em auth.users.email. Zero ou multiplas contas abortam.
-- Se o identificador nao for esse, substitua o resolvedor por UUIDs conferidos.
begin;
create temporary table import_players (
  name text not null, account text primary key, player_id uuid unique
) on commit drop;
insert into import_players(name, account) values
""" + ',\n'.join(f'({quote(n)}, {quote(a)})' for n, a in players) + ';\n'
sql += """
-- Validar todas as identidades antes de qualquer escrita persistente.
do $$
declare p record; matches integer; resolved uuid;
begin
  if not exists (select 1 from public.tournaments
    where id = '20260900-0000-4000-8000-000000000001') then
    raise exception 'Execute primeiro o seed september.sql';
  end if;
  for p in select * from import_players loop
    select count(*) into matches from auth.users u
      join public.profiles pr on pr.id = u.id
      where split_part(u.email, '@', 1) = p.account;
    if matches <> 1 then
      raise exception 'Conta %: % correspondencias. Confira o UUID antes de carregar.', p.account, matches;
    end if;
    select u.id into resolved from auth.users u
      join public.profiles pr on pr.id = u.id
      where split_part(u.email, '@', 1) = p.account;
    update import_players set player_id = resolved where account = p.account;
  end loop;
end $$;

create temporary table import_scores (
  account text references import_players(account), game_day date, raw_score integer,
  primary key(account, game_day), check (raw_score between 0 and 25000)
) on commit drop;
insert into import_scores values
""" + ',\n'.join(f'({quote(a)}, {quote(d)}, {v})' for a, d, v in scores) + ';\n'
sql += """
do $$
declare player_row record; score_row record;
begin
  for player_row in select * from import_players order by account loop
    perform private.add_september_participant(player_row.player_id, player_row.name);
  end loop;
  for score_row in select p.player_id, s.game_day, s.raw_score
    from import_scores s join import_players p using(account)
    order by s.game_day, p.account loop
    perform private.load_september_score(score_row.player_id, score_row.game_day, score_row.raw_score);
  end loop;
  if exists (
    select 1 from import_scores s join import_players p using(account)
    left join public.personal_scores actual
      on actual.player_id = p.player_id and actual.played_on = s.game_day
    where actual.id is null or actual.score <> s.raw_score
  ) then
    raise exception 'Falha na conciliacao da carga';
  end if;
end $$;

-- Resumo da carga conferida. Lacunas nao foram convertidas em zero.
select p.name, p.account, p.player_id, count(s.game_day) as scores_imported
from import_players p left join import_scores s using(account)
group by p.name, p.account, p.player_id order by p.name;
commit;
"""
target = Path(__file__).with_name('september-history-2026-09-22.sql')
target.write_text(sql, encoding='utf-8')
print(f'{len(players)} participantes; {len(scores)} pontuacoes; pendentes: {pending}')
