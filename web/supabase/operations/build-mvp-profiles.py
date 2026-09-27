"""Read the unchanged MVP workbook and generate the one-off profile import."""
from hashlib import sha256
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[3]
source = root / 'GeoGuaras.xlsx'
digest = sha256(source.read_bytes()).hexdigest()
assert digest == 'e74d1a6560a56ee1f7d3103df641199fbf308b2a5e735a46b6e0082bf93efcc4'
sheet = openpyxl.load_workbook(source)['Diario']
ids = {'Luca': '8d899496-3cb0-41f6-b67d-4e57bb1185c4',
       'Zade': '55a4aeaf-30bc-403a-a5e6-563328466313'}
def quote(value):
    return 'null' if value is None else "'" + str(value).replace("'", "''") + "'"
rows = []
for row in range(2, 17):
    legacy, cell, account = sheet.cell(row, 1).value, sheet.cell(row, 2), sheet.cell(row, 3).value
    # User already filled their own profile; preserve Martin entirely.
    if legacy == 'Martin':
        continue
    if cell.value is None:
        assert legacy in ('Ramiro', 'Marcelo') and account is None
        continue
    assert account or legacy in ids
    rows.append((legacy, account, ids.get(legacy), cell.value,
                 cell.hyperlink.target if cell.hyperlink else None))
assert len(rows) == 12 and sum(r[4] is not None for r in rows) == 11
values = ',\n'.join('(' + ', '.join(map(quote, row)) + ')' for row in rows)
sql = f"""-- Carga pontual autorizada pelo Dono do produto em 27/09/2026.
-- Fonte: GeoGuaras.xlsx / Diario!A2:C16 (nome publico e hyperlink da coluna B).
-- SHA256: {digest}
-- Identidades: coluna C confirmada na carga MVP; Luca/Zade usam UUIDs confirmados.
-- Executar INTEGRALMENTE como administrador no SQL Editor, apos S4-04/S4-05.
-- 12 nomes, 11 URLs; Martin fora da carga: perfil preenchido pelo titular.
-- Luca sem URL na fonte: preserva a URL atual.
-- Valores manuais divergentes ou contas ausentes/ambiguas abortam tudo.
-- Auditoria privada existente guarda o antes/depois; nao altera o backup original.
begin;
lock table auth.users in share mode;
lock table public.profiles in share row exclusive mode;
create temporary table mvp_profile_import (
  legacy_name text primary key, account text unique, player_id uuid unique,
  public_name text not null, profile_url text
) on commit drop;
insert into mvp_profile_import values
{values};

do $$
declare item record; matches integer; resolved uuid; current_profile public.profiles;
begin
  for item in select * from mvp_profile_import order by legacy_name loop
    if item.player_id is null then
      select count(*), min(u.id::text)::uuid into matches, resolved
      from auth.users u join public.profiles p on p.id=u.id
      where split_part(u.email,'@',1)=item.account;
    else
      select count(*), min(u.id::text)::uuid into matches, resolved
      from auth.users u join public.profiles p on p.id=u.id where u.id=item.player_id;
    end if;
    if matches <> 1 then
      raise exception 'mvp_profile_identity: % tem % correspondencias', item.legacy_name, matches;
    end if;
    update mvp_profile_import set player_id=resolved,
      public_name=private.validate_public_name(item.public_name),
      profile_url=private.clean_geoguessr_url(item.profile_url)
      where legacy_name=item.legacy_name;
    select * into strict current_profile from public.profiles where id=resolved;
    if current_profile.display_name is distinct from item.legacy_name
       and current_profile.display_name is distinct from item.public_name then
      raise exception 'mvp_profile_name_changed: %; conferir nome atual antes da carga', item.legacy_name;
    end if;
    if item.profile_url is not null and current_profile.geoguessr_url is not null
       and current_profile.geoguessr_url is distinct from private.clean_geoguessr_url(item.profile_url) then
      raise exception 'mvp_profile_url_changed: %; conferir URL atual antes da carga', item.legacy_name;
    end if;
  end loop;
  if exists(select 1 from mvp_profile_import i join public.profiles p
    on private.public_name_key(p.display_name)=private.public_name_key(i.public_name)
    and p.id<>i.player_id) then
    raise exception 'mvp_profile_name_unavailable: nome da planilha reservado por outra conta';
  end if;
end $$;

-- A trigger private.log_profile_change registra somente perfis efetivamente alterados.
update public.profiles p set display_name=i.public_name, public_name_confirmed=true,
  geoguessr_url=coalesce(i.profile_url,p.geoguessr_url)
from mvp_profile_import i where p.id=i.player_id
  and (p.display_name,p.public_name_confirmed,p.geoguessr_url)
    is distinct from (i.public_name,true,coalesce(i.profile_url,p.geoguessr_url));

do $$ begin
  if (select count(*) from public.profiles p join mvp_profile_import i on i.player_id=p.id
    where p.display_name=i.public_name and p.public_name_confirmed
      and (i.profile_url is null or p.geoguessr_url=i.profile_url)) <> 12 then
    raise exception 'mvp_profile_reconciliation_failed';
  end if;
end $$;

select i.legacy_name as nome_mvp, p.id, p.display_name as nome_publico,
  p.geoguessr_url, p.public_name_confirmed as confirmado,
  case when i.profile_url is null then 'Sem URL na planilha; valor atual preservado'
       else 'Nome e URL conferidos' end as conferencia
from mvp_profile_import i join public.profiles p on p.id=i.player_id order by i.legacy_name;
commit;
"""
Path(__file__).with_name('20260927-import-mvp-profiles.sql').write_text(sql, encoding='utf-8')
print('SQL gerado: 12 nomes e 11 URLs; Martin preservado; fonte intacta.')
